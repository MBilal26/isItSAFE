import os
import subprocess
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import threading
from PIL import Image
from PyPDF2 import PdfReader, PdfWriter
from docx import Document

# Logger integration
import sys
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
try:
    from logger import log_event
except ImportError:
    def log_event(m, msg, l="info"): pass

# ============================================================================
#                          CONFIGURATION
# ============================================================================
BG_DARK = "#0B0E14"
BG_SIDE = "#151921"
BG_CARD = "#1E293B"
ACCENT = "#00D1FF"
DANGER = "#FF3131"
SUCCESS = "#00FF41"
WARNING = "#FFB800"
TEXT_MAIN = "#E0E6ED"
TEXT_DIM = "#888888"

OUTPUT_DIR = "output_files"
os.makedirs(OUTPUT_DIR, exist_ok=True)

class MetadataCleanerApp:
    def __init__(self, root):
        self.root = root
        self.root.title("isItSafe - Metadata Privacy Cleaner")
        self.root.state('zoomed')
        self.root.configure(bg=BG_DARK)
        
        self.selected_file = ""
        self.before_meta = {}
        self.is_processing = False
        
        self._setup_sidebar()
        self._setup_main_area()
        self._setup_styles()

    def _setup_styles(self):
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("TScrollbar", gripcount=0, background=BG_CARD, darkcolor=BG_DARK, lightcolor=BG_DARK, bordercolor=BG_DARK, arrowcolor=ACCENT)

    def _setup_sidebar(self):
        self.sidebar = tk.Frame(self.root, bg=BG_SIDE, width=280)
        self.sidebar.pack(side=tk.LEFT, fill=tk.Y)
        self.sidebar.pack_propagate(False)
        
        # Logo
        tk.Label(self.sidebar, text="isItSafe", font=("Consolas", 28, "bold"), bg=BG_SIDE, fg=ACCENT).pack(pady=(40, 10))
        tk.Label(self.sidebar, text="Metadata Cleaner", font=("Segoe UI", 9), bg=BG_SIDE, fg=TEXT_DIM).pack()
        
        self.canvas = tk.Canvas(self.sidebar, width=150, height=150, bg=BG_SIDE, highlightthickness=0)
        self.canvas.pack(pady=30)
        self.canvas.create_text(75, 75, text="🖼️", font=("Segoe UI", 80), fill=BG_CARD)
        
        # Status Label
        self.status_var = tk.StringVar(value="READY")
        self.status_label = tk.Label(self.sidebar, textvariable=self.status_var, font=("Segoe UI", 10, "bold"), bg=BG_SIDE, fg=TEXT_DIM)
        self.status_label.pack(pady=10)

        # Actions
        btn_frame = tk.Frame(self.sidebar, bg=BG_SIDE)
        btn_frame.pack(fill=tk.X, padx=20, pady=10)
        
        self.select_btn = tk.Button(
            btn_frame, text="📁 SELECT FILE", command=self.select_file,
            bg=BG_CARD, fg=TEXT_MAIN, font=("Segoe UI", 10, "bold"),
            relief="flat", cursor="hand2", pady=12
        )
        self.select_btn.pack(fill=tk.X, pady=5)
        
        self.clean_btn = tk.Button(
            btn_frame, text="✨ CLEAN METADATA", command=self.clean_metadata,
            bg=SUCCESS, fg=BG_DARK, font=("Segoe UI", 10, "bold"),
            relief="flat", cursor="hand2", pady=12, state="disabled"
        )
        self.clean_btn.pack(fill=tk.X, pady=5)

        tk.Label(
            self.sidebar, text="Supports: JPG, PNG, PDF, DOCX",
            font=("Segoe UI", 8), bg=BG_SIDE, fg=TEXT_DIM
        ).pack(side=tk.BOTTOM, pady=20)

    def _setup_main_area(self):
        self.main = tk.Frame(self.root, bg=BG_DARK)
        self.main.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=30, pady=30)
        
        # Header
        tk.Label(
            self.main, text="🛡️ Metadata Privacy Shield",
            font=("Segoe UI", 22, "bold"), bg=BG_DARK, fg=TEXT_MAIN
        ).pack(anchor=tk.W, pady=(0, 20))
        
        # File Info Card
        self.info_card = tk.Frame(self.main, bg=BG_CARD, padx=20, pady=20)
        self.info_card.pack(fill=tk.X, pady=(0, 20))
        
        self.file_label = tk.Label(
            self.info_card, text="PLEASE SELECT A FILE TO BEGIN ANALYSIS",
            font=("Consolas", 10), bg=BG_CARD, fg=TEXT_DIM, wraplength=800
        )
        self.file_label.pack()

        # Tabs/Buttons for metadata view
        tab_frame = tk.Frame(self.main, bg=BG_DARK)
        tab_frame.pack(fill=tk.X)
        
        self.view_before_btn = tk.Button(
            tab_frame, text="ANALYZED METADATA", command=lambda: self.view_metadata(True),
            bg=BG_CARD, fg=WARNING, font=("Segoe UI", 9, "bold"),
            relief="flat", padx=25, pady=10, state="disabled"
        )
        self.view_before_btn.pack(side=tk.LEFT, padx=(0, 5))
        
        self.view_after_btn = tk.Button(
            tab_frame, text="PREVIEW CLEANED STATE", command=lambda: self._run_preview_task(),
            bg=BG_CARD, fg=SUCCESS, font=("Segoe UI", 9, "bold"),
            relief="flat", padx=25, pady=10, state="disabled"
        )
        self.view_after_btn.pack(side=tk.LEFT)

        # Content Area (Scrollable Text)
        self.text_frame = tk.Frame(self.main, bg=BG_CARD, pady=2)
        self.text_frame.pack(fill=tk.BOTH, expand=True, pady=10)
        
        self.scrollbar = ttk.Scrollbar(self.text_frame)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.text_box = tk.Text(
            self.text_frame, bg=BG_DARK, fg=TEXT_MAIN, font=("Consolas", 10),
            padx=20, pady=20, relief="flat", yscrollcommand=self.scrollbar.set,
            wrap=tk.WORD, undo=True
        )
        self.text_box.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.scrollbar.config(command=self.text_box.yview)

    def set_busy(self, status="PROCESSING..."):
        self.is_processing = True
        self.status_var.set(status)
        self.status_label.config(fg=ACCENT)
        self.select_btn.config(state="disabled")
        self.clean_btn.config(state="disabled")
        self.view_before_btn.config(state="disabled")
        self.view_after_btn.config(state="disabled")
        self.root.config(cursor="watch")

    def set_ready(self, status="READY"):
        self.is_processing = False
        self.status_var.set(status)
        self.status_label.config(fg=TEXT_DIM)
        self.select_btn.config(state="normal")
        if self.selected_file:
            self.clean_btn.config(state="normal")
            self.view_before_btn.config(state="normal")
            self.view_after_btn.config(state="normal")
        self.root.config(cursor="")

    def select_file(self):
        file_path = filedialog.askopenfilename(
            title="Select a File",
            filetypes=[("Supported Files", "*.jpg *.jpeg *.png *.pdf *.docx")]
        )
        if file_path:
            self.selected_file = file_path
            self.file_label.config(text=f"TARGET: {os.path.basename(self.selected_file)}", fg=ACCENT)
            log_event("Metadata_PIC", f"File selected for analysis: {os.path.basename(file_path)}")
            self.set_busy("ANALYZING...")
            threading.Thread(target=self._analyze_file_task, daemon=True).start()

    def _analyze_file_task(self):
        try:
            self.before_meta = self.get_full_metadata(self.selected_file)
            self.root.after(0, lambda: self.view_metadata(True))
            self.root.after(0, lambda: self.set_ready("FILE LOADED"))
        except Exception as e:
            self.root.after(0, lambda: messagebox.showerror("Analysis Error", str(e)))
            self.root.after(0, self.set_ready)

    def get_full_metadata(self, file_path):
        meta = self._get_native_metadata(file_path)
        try:
            result = subprocess.run(
                ["exiftool", file_path], stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                text=True, encoding="utf-8", errors="ignore"
            )
            if result.returncode == 0:
                lines = result.stdout.strip().split("\n")
                exif_meta = {line.split(":", 1)[0].strip(): line.split(":", 1)[1].strip() for line in lines if ":" in line}
                meta.update(exif_meta)
        except: pass
        return meta if meta else {"Status": "No metadata found or format unsupported"}

    def _get_native_metadata(self, file_path):
        ext = os.path.splitext(file_path)[1].lower()
        meta = {"Internal Filename": os.path.basename(file_path)}
        try:
            if ext in [".jpg", ".jpeg", ".png"]:
                with Image.open(file_path) as img:
                    meta["Dimensions"] = f"{img.width}x{img.height}"
                    meta["Format"] = img.format
                    meta["Mode"] = img.mode
                    if hasattr(img, '_getexif') and img._getexif():
                        for tag, value in img._getexif().items():
                            meta[f"EXIF_{tag}"] = str(value)
            elif ext == ".pdf":
                reader = PdfReader(file_path)
                if reader.metadata:
                    for k, v in reader.metadata.items():
                        meta[f"PDF_{k[1:] if k.startswith('/') else k}"] = str(v)
                meta["Pages"] = str(len(reader.pages))
            elif ext == ".docx":
                doc = Document(file_path)
                props = doc.core_properties
                attrs = ['author', 'category', 'comments', 'content_status', 'created', 'identifier', 
                        'keywords', 'language', 'last_modified_by', 'last_printed', 'modified', 
                        'revision', 'subject', 'title', 'version']
                for attr in attrs:
                    val = getattr(props, attr)
                    if val: meta[f"DOCX_{attr.title()}"] = str(val)
        except Exception as e:
            meta["Native Analysis Error"] = str(e)
        return meta

    def view_metadata(self, before=True):
        self.text_box.config(state="normal")
        self.text_box.delete("1.0", tk.END)
        if before:
            content = "🔍 DETECTED METADATA ATRIBUTES:\n" + "="*50 + "\n\n"
            content += "\n".join([f"{k:35}: {v}" for k, v in self.before_meta.items()])
            self.text_box.insert(tk.END, content)
        self.text_box.config(state="disabled")

    def _run_preview_task(self):
        self.set_busy("GENERATING PREVIEW...")
        threading.Thread(target=self._preview_task, daemon=True).start()

    def _preview_task(self):
        try:
            temp_path = self.clean_file_temp(self.selected_file)
            after_meta = self.get_full_metadata(temp_path)
            if os.path.exists(temp_path): os.remove(temp_path)
            removed = {k: v for k, v in self.before_meta.items() if k not in after_meta}
            changed = {k: (self.before_meta[k], after_meta[k]) for k in self.before_meta if k in after_meta and self.before_meta[k] != after_meta[k]}
            self.root.after(0, lambda: self._update_preview_ui(removed, changed))
        except Exception as e:
            self.root.after(0, lambda: messagebox.showerror("Preview Error", str(e)))
        finally:
            self.root.after(0, self.set_ready)

    def _update_preview_ui(self, removed, changed):
        self.text_box.config(state="normal")
        self.text_box.delete("1.0", tk.END)
        self.text_box.insert(tk.END, "✨ PRIVACY IMPACT PREVIEW\n", "header")
        self.text_box.insert(tk.END, "="*50 + "\n\n")
        self.text_box.insert(tk.END, "🚫 REMOVED DATA:\n", "header_rem")
        if not removed: self.text_box.insert(tk.END, "None (File might already be clean)\n")
        for k, v in removed.items(): self.text_box.insert(tk.END, f"{k:35}: {v}\n", "rem")
        self.text_box.insert(tk.END, "\n♻️ MODIFIED/SANITIZED DATA:\n", "header_chg")
        if not changed: self.text_box.insert(tk.END, "None\n")
        for k, (b, a) in changed.items(): self.text_box.insert(tk.END, f"{k:35}: {b} ➡️ {a}\n", "chg")
        self.text_box.tag_config("header", foreground=ACCENT, font=("Consolas", 14, "bold"))
        self.text_box.tag_config("header_rem", foreground=DANGER, font=("Consolas", 11, "bold"))
        self.text_box.tag_config("rem", foreground="#FF7A7A")
        self.text_box.tag_config("header_chg", foreground=WARNING, font=("Consolas", 11, "bold"))
        self.text_box.tag_config("chg", foreground="#FFD67A")
        self.text_box.config(state="disabled")

    def clean_file_temp(self, file_path):
        ext = os.path.splitext(file_path)[1].lower()
        temp_path = os.path.join(OUTPUT_DIR, "temp_preview" + ext)
        self._process_cleaning(file_path, temp_path)
        return temp_path

    def clean_metadata(self):
        self.set_busy("CLEANING FILE...")
        threading.Thread(target=self._clean_task, daemon=True).start()

    def _clean_task(self):
        try:
            output_path = os.path.join(OUTPUT_DIR, "Cleaned_" + os.path.basename(self.selected_file))
            self._process_cleaning(self.selected_file, output_path)
            self.root.after(0, lambda: messagebox.showinfo("Success", f"Privacy Shield Applied!\n\nSaved to: {output_path}"))
        except Exception as e:
            self.root.after(0, lambda: messagebox.showerror("Cleaning Error", str(e)))
        finally:
            self.root.after(0, self.set_ready)

    def _process_cleaning(self, source, dest):
        ext = os.path.splitext(source)[1].lower()
        if ext in [".jpg", ".jpeg", ".png"]:
            with Image.open(source) as img:
                data = list(img.getdata())
                clean_img = Image.new(img.mode, img.size)
                clean_img.putdata(data)
                clean_img.save(dest, quality=95, optimize=True)
        elif ext == ".pdf":
            reader = PdfReader(source); writer = PdfWriter(); [writer.add_page(p) for p in reader.pages]
            with open(dest, "wb") as f: writer.write(f)
        elif ext == ".docx":
            doc = Document(source)
            props = doc.core_properties
            attrs = ['author', 'category', 'comments', 'content_status', 'created', 'identifier', 'keywords', 'language', 'last_modified_by', 'last_printed', 'modified', 'revision', 'subject', 'title', 'version']
            for attr in attrs: setattr(props, attr, None)
            doc.save(dest)

if __name__ == "__main__":
    root = tk.Tk(); app = MetadataCleanerApp(root); root.mainloop()